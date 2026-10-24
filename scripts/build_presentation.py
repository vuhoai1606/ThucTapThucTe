# -*- coding: utf-8 -*-
r"""
Script: build_presentation.py
Author: Antigravity
Purpose:
  Build a high-impact, professional, 16-slide English PowerPoint presentation 
  showcasing the FSOFT HRMS product developed during the Practical Internship 
  at FPT Software Da Nang by student Thai Hoai Vu (24IT314).
Output file:
  D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\Slide-TTTT-ThaiHoaiVu-24IT314.pptx
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = r"D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\Slide-TTTT-ThaiHoaiVu-24IT314.pptx"
SCREENSHOTS_DIR = r"D:\Lap Trinh\Nam 3 Ki 1\ThucTapThucTe\screenshots"
EXTRACTED_IMG_DIR = r"D:\Lap Trinh\Nam 3 Ki 1\ThucTapThucTe\extracted_doc_images"

# --- COLOR PALETTE ---
NAVY_DARK     = RGBColor(11, 27, 61)     # #0B1B3D - Hero Background
NAVY_CARD     = RGBColor(16, 42, 84)     # #102A54 - Dark Card Background
BG_LIGHT      = RGBColor(248, 250, 252)  # #F8FAFC - Clean Light Slate
WHITE         = RGBColor(255, 255, 255)  # #FFFFFF - Pure White
CARD_BORDER   = RGBColor(226, 232, 240)  # #E2E8F0 - Subtle Border
TEXT_DARK     = RGBColor(15, 23, 42)     # #0F172A - High Contrast Dark Text
TEXT_MUTED    = RGBColor(100, 116, 139)  # #64748B - Slate Subtitle
TEXT_LIGHT    = RGBColor(241, 245, 249)  # #F1F5F9 - Light Text
BRAND_BLUE    = RGBColor(0, 102, 204)    # #0066CC - Tech Blue
BRAND_ORANGE  = RGBColor(242, 112, 36)   # #F27024 - FPT Orange Accent
ACCENT_CYAN   = RGBColor(14, 165, 233)   # #0EA5E9 - Cyan Highlight
SUCCESS_GREEN = RGBColor(16, 185, 129)   # #10B981 - Success Green
TAG_BG        = RGBColor(238, 242, 255)  # #EEF2FF - Light Tag BG
TAG_TEXT      = RGBColor(67, 56, 202)    # #4338CA - Tag Text

TOTAL_SLIDES = 16

def get_screenshot(filename):
    p = os.path.join(SCREENSHOTS_DIR, filename)
    if os.path.exists(p):
        return p
    return None

def get_extracted_img(filename):
    p = os.path.join(EXTRACTED_IMG_DIR, filename)
    if os.path.exists(p):
        return p
    return None

def set_shape_border(shape, color, width_pt=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)

def add_header(slide, tag_text, title_text, subtitle_text, is_dark=False):
    """Adds a standard structured header to a content slide."""
    # 1. Tag pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(8), Inches(0.28))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.name = "Segoe UI"
    p_tag.font.color.rgb = BRAND_ORANGE if not is_dark else ACCENT_CYAN

    # 2. Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.70), Inches(11.7), Inches(0.48))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.name = "Segoe UI"
    p_title.font.color.rgb = WHITE if is_dark else TEXT_DARK

    # 3. Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.18), Inches(11.7), Inches(0.35))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(11)
    p_sub.font.name = "Segoe UI"
    p_sub.font.color.rgb = RGBColor(148, 163, 184) if is_dark else TEXT_MUTED

def add_footer(slide, current_idx, is_dark=False):
    """Adds a standard clean footer with divider line, presenter details, and page numbering."""
    # Divider line
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.015))
    div.fill.solid()
    div.fill.fore_color.rgb = RGBColor(30, 58, 110) if is_dark else CARD_BORDER
    div.line.fill.background()

    # Left footer text
    box_left = slide.shapes.add_textbox(Inches(0.8), Inches(6.98), Inches(9.0), Inches(0.3))
    tf_l = box_left.text_frame
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "Thai Hoai Vu (24IT314) • Practical Internship Capstone • FPT Software Da Nang"
    p_l.font.size = Pt(9)
    p_l.font.name = "Segoe UI"
    p_l.font.color.rgb = RGBColor(148, 163, 184) if is_dark else RGBColor(148, 163, 184)

    # Right footer text
    box_right = slide.shapes.add_textbox(Inches(10.0), Inches(6.98), Inches(2.533), Inches(0.3))
    tf_r = box_right.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = f"Slide {current_idx} of {TOTAL_SLIDES}"
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.font.size = Pt(9)
    p_r.font.bold = True
    p_r.font.name = "Segoe UI"
    p_r.font.color.rgb = ACCENT_CYAN if is_dark else BRAND_BLUE

def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER):
    """Adds a rounded rectangle container card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    set_shape_border(card, border_color, 1)
    return card

def add_feature_bullets(text_frame, items, is_dark=False):
    """Adds styled bullet points with bold titles and descriptions."""
    for i, (b_title, desc) in enumerate(items):
        p = text_frame.add_paragraph() if i > 0 or len(text_frame.paragraphs[0].text) > 0 else text_frame.paragraphs[0]
        p.space_after = Pt(8)
        p.font.name = "Segoe UI"
        
        # Bullet indicator / Title
        r_title = p.add_run()
        r_title.text = f"•  {b_title}: "
        r_title.font.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = WHITE if is_dark else TEXT_DARK
        
        # Description
        r_desc = p.add_run()
        r_desc.text = desc
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = RGBColor(203, 213, 225) if is_dark else RGBColor(51, 65, 85)

def add_tech_tags(slide, left, top, tags):
    """Adds horizontal tech badge pills."""
    curr_x = left
    for tag in tags:
        tag_width = Inches(max(0.7, len(tag) * 0.085 + 0.3))
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, curr_x, top, tag_width, Inches(0.24))
        pill.fill.solid()
        pill.fill.fore_color.rgb = TAG_BG
        set_shape_border(pill, RGBColor(199, 210, 254), 0.75)
        
        tf = pill.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = tag
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.name = "Segoe UI"
        p.font.color.rgb = TAG_TEXT
        
        curr_x += tag_width + Inches(0.12)

def add_screenshot_card(slide, img_path, left, top, width, height, caption=None):
    """Embeds an image inside a neat white frame card."""
    card = add_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER)
    
    if img_path and os.path.exists(img_path):
        # Calculate image position with padding inside card
        pad = Inches(0.12)
        target_w = width - (pad * 2)
        # 16:10 screenshot ratio: height = target_w / 1.6
        calc_h = target_w / 1.6
        if calc_h > (height - Inches(0.5)):
            calc_h = height - Inches(0.5)
            target_w = calc_h * 1.6
            
        img_x = left + (width - target_w) / 2
        img_y = top + pad + Inches(0.05)
        
        pic = slide.shapes.add_picture(img_path, img_x, img_y, width=target_w, height=calc_h)
        set_shape_border(pic, RGBColor(203, 213, 225), 0.5)
        
        if caption:
            cap_box = slide.shapes.add_textbox(left, top + height - Inches(0.32), width, Inches(0.26))
            tf = cap_box.text_frame
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = caption
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(9)
            p.font.name = "Segoe UI"
            p.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE BUILDERS
# ==============================================================================

def build_slide_1_title(prs):
    """Slide 1: Hero Title Slide (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    # Dark Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DARK
    bg.line.fill.background()

    # Subtle top accent bar (FPT Orange + Tech Blue)
    bar1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.5), Inches(0.08))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = BRAND_ORANGE
    bar1.line.fill.background()
    bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(0), Inches(8.833), Inches(0.08))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = BRAND_BLUE
    bar2.line.fill.background()

    # Logos (VKU & FPT Software)
    vku_logo = get_extracted_img("image1.jpg")
    fpt_logo = get_extracted_img("image2.png")
    if vku_logo:
        slide.shapes.add_picture(vku_logo, Inches(0.8), Inches(0.45), height=Inches(0.65))
    if fpt_logo:
        slide.shapes.add_picture(fpt_logo, Inches(2.2), Inches(0.45), height=Inches(0.65))

    # University header text
    hdr_box = slide.shapes.add_textbox(Inches(3.3), Inches(0.48), Inches(9.0), Inches(0.6))
    tf_h = hdr_box.text_frame
    tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0
    p1 = tf_h.paragraphs[0]
    p1.text = "VIETNAM - KOREA UNIVERSITY OF INFORMATION AND COMMUNICATION TECHNOLOGY (VKU)"
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(148, 163, 184)
    p2 = tf_h.add_paragraph()
    p2.text = "FACULTY OF COMPUTER SCIENCE • PRACTICAL INTERNSHIP CAPSTONE REPORT"
    p2.font.size = Pt(9)
    p2.font.color.rgb = ACCENT_CYAN

    # Main Project Title Box
    title_card = add_card(slide, Inches(0.8), Inches(1.5), Inches(11.733), Inches(2.7), bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    
    t_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.75), Inches(10.9), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p_tag = tf.paragraphs[0]
    p_tag.text = "ENTERPRISE SOFTWARE PRODUCT PRESENTATION"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = BRAND_ORANGE
    
    p_main = tf.add_paragraph()
    p_main.text = "FSOFT HRMS"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE
    p_main.space_before = Pt(6)
    
    p_sub = tf.add_paragraph()
    p_sub.text = "Human Resource & Enterprise Workforce Allocation Management System"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(226, 232, 240)
    p_sub.space_before = Pt(4)
    
    p_desc = tf.add_paragraph()
    p_desc.text = "A comprehensive Single Page Application (SPA) designed to digitize workforce operations, project staffing allocation, leave approval workflows, automated payroll, and KPI performance appraisals."
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = RGBColor(148, 163, 184)
    p_desc.space_before = Pt(8)

    # 3 Information Meta Cards
    col_w = Inches(3.75)
    gap = Inches(0.24)
    
    # Card 1: Student Presenter
    c1 = add_card(slide, Inches(0.8), Inches(4.45), col_w, Inches(2.0), bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(4.6), col_w - Inches(0.4), Inches(1.7))
    tf1 = tb1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "STUDENT PRESENTER"
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    p2 = tf1.add_paragraph()
    p2.text = "Thái Hoài Vũ (Thai Hoai Vu)"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)
    p3 = tf1.add_paragraph()
    p3.text = "Student ID: 24IT314 • Class: 24JIT\nFaculty of Computer Science, VKU"
    p3.font.size = Pt(10)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(4)

    # Card 2: Academic Advisor
    c2 = add_card(slide, Inches(0.8) + col_w + gap, Inches(4.45), col_w, Inches(2.0), bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb2 = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, Inches(4.6), col_w - Inches(0.4), Inches(1.7))
    tf2 = tb2.text_frame
    p1 = tf2.paragraphs[0]
    p1.text = "ACADEMIC ADVISOR"
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    p2 = tf2.add_paragraph()
    p2.text = "M.Sc. Nguyễn Ngọc Huyền Trân"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)
    p3 = tf2.add_paragraph()
    p3.text = "Lecturer, Faculty of Computer Science\nVietnam - Korea University (VKU)"
    p3.font.size = Pt(10)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(4)

    # Card 3: Host Enterprise
    c3 = add_card(slide, Inches(0.8) + (col_w + gap)*2, Inches(4.45), col_w, Inches(2.0), bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb3 = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, Inches(4.6), col_w - Inches(0.4), Inches(1.7))
    tf3 = tb3.text_frame
    p1 = tf3.paragraphs[0]
    p1.text = "HOST ENTERPRISE"
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = BRAND_ORANGE
    p2 = tf3.add_paragraph()
    p2.text = "FPT Software Đà Nẵng"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)
    p3 = tf3.add_paragraph()
    p3.text = "Campus: FPT Complex Da Nang\nCMMI Level 5 Software Development"
    p3.font.size = Pt(10)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(4)

    # Bottom Footer on Title
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Tech Stack: Node.js • Express.js • PostgreSQL • RESTful API • Role-Based Access Control (RBAC) • JWT • Bcrypt"
    p_f.alignment = PP_ALIGN.CENTER
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = RGBColor(148, 163, 184)

def build_slide_2_motivation(prs):
    """Slide 2: Project Motivation & Industry Challenges"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    add_header(slide, "1. Motivation & Problem Statement", "Enterprise HR & Resource Allocation Challenges", "Practical observations and pain points identified during internship at FPT Software Da Nang.")
    add_footer(slide, 2)
    
    # 3 Cards for the 3 Major Pain Points
    card_w = Inches(3.72)
    gap = Inches(0.28)
    top_y = Inches(1.75)
    h_y = Inches(4.9)

    # Card 1: Data Fragmentation
    add_card(slide, Inches(0.8), top_y, card_w, h_y)
    tb = slide.shapes.add_textbox(Inches(1.0), top_y + Inches(0.2), card_w - Inches(0.4), h_y - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Data Fragmentation"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    
    items1 = [
        ("Disparate Excel Files", "Employee records, skills, and project logs maintained across isolated spreadsheets."),
        ("Information Silos", "Communication barriers between Project Managers (PMs) and Human Resource (HR) leads."),
        ("Outdated Profiles", "Developer technical skills and completed certifications not updated in real time."),
        ("Reporting Friction", "Compiling cross-departmental headcount reports required tedious manual consolidation.")
    ]
    add_feature_bullets(tf, items1)

    # Card 2: Staffing & Bench Inefficiency (Key Breakthrough)
    add_card(slide, Inches(0.8) + card_w + gap, top_y, card_w, h_y, border_color=BRAND_ORANGE)
    tb2 = slide.shapes.add_textbox(Inches(1.0) + card_w + gap, top_y + Inches(0.2), card_w - Inches(0.4), h_y - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "2. Resource Staffing Bottlenecks"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = BRAND_ORANGE
    
    items2 = [
        ("Lack of Visibility", "No single-pane view showing exact developer workload allocation percentages."),
        ("Burnout Risk", "Key technical leads often assigned to multiple projects exceeding 100% capacity."),
        ("Idle Bench Overhead", "Engineers on bench (unassigned) were difficult to identify and staff promptly."),
        ("Skill Misalignment", "Project demands mismatched with developer core competences, slowing delivery.")
    ]
    add_feature_bullets(tf2, items2)

    # Card 3: Administrative Lag
    add_card(slide, Inches(0.8) + (card_w + gap)*2, top_y, card_w, h_y)
    tb3 = slide.shapes.add_textbox(Inches(1.0) + (card_w + gap)*2, top_y + Inches(0.2), card_w - Inches(0.4), h_y - Inches(0.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "3. Administrative Overhead"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = BRAND_BLUE
    
    items3 = [
        ("Manual Leave Routing", "Leave approvals routed via paper or scattered email threads causing delays."),
        ("Complex Payroll", "High risk of human error calculating base salary, overtime, and allowances."),
        ("Subjective Appraisals", "Performance KPI reviews lacked systematic rubric scoring and multi-tier signoff."),
        ("Security Compliance", "Absence of strict Role-Based Access Control risking confidential data exposure.")
    ]
    add_feature_bullets(tf3, items3)

def build_slide_3_architecture(prs):
    """Slide 3: System Architecture & Modern Tech Stack"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    add_header(slide, "2. System Architecture & Tech Stack", "Robust 3-Tier Enterprise Fullstack Architecture", "Engineered for high performance, modularity, strict security compliance, and sub-second API response.")
    add_footer(slide, 3)
    
    col_w = Inches(3.72)
    gap = Inches(0.28)
    top_y = Inches(1.75)
    h_y = Inches(4.9)

    # Tier 1: Client Layer
    add_card(slide, Inches(0.8), top_y, col_w, h_y)
    tb = slide.shapes.add_textbox(Inches(1.0), top_y + Inches(0.2), col_w - Inches(0.4), h_y - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Presentation Layer (Client)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    
    items1 = [
        ("SPA Architecture", "Single Page Application delivering fluid, seamless UI transitions with zero reloads."),
        ("Modern Semantic Web", "HTML5 & CSS3 with modular, responsive layouts for all desktop displays."),
        ("Vanilla ES6+ JavaScript", "Ultra-lightweight footprint with zero heavy framework bloat, maximizing client speed."),
        ("Interactive Chart.js", "Dynamic charts for FSU headcount, project progress radar, and KPI distributions."),
        ("Route Guard Protection", "Client-side state guards preventing unauthorized URL routing.")
    ]
    add_feature_bullets(tf, items1)
    add_tech_tags(slide, Inches(1.0), top_y + h_y - Inches(0.55), ["SPA", "HTML5/CSS3", "ES6+", "Chart.js"])

    # Tier 2: Server Layer
    add_card(slide, Inches(0.8) + col_w + gap, top_y, col_w, h_y, border_color=BRAND_ORANGE)
    tb2 = slide.shapes.add_textbox(Inches(1.0) + col_w + gap, top_y + Inches(0.2), col_w - Inches(0.4), h_y - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Business Logic Layer (Server)"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = BRAND_ORANGE
    
    items2 = [
        ("Node.js Runtime", "Event-driven, non-blocking asynchronous I/O handling high concurrency efficiently."),
        ("Express.js Framework", "Robust RESTful API architecture with standardized HTTP methods (GET, POST, PUT, DELETE)."),
        ("Strict MVC Pattern", "Modular separation of concerns: Routes, Controllers, Models, and Middlewares."),
        ("Business Validation", "Server-side data validation ensuring payload integrity before database writes."),
        ("CORS & Error Pipeline", "Centralized error handling and cross-origin resource security configuration.")
    ]
    add_feature_bullets(tf2, items2)
    add_tech_tags(slide, Inches(1.0) + col_w + gap, top_y + h_y - Inches(0.55), ["Node.js", "Express", "REST API", "MVC"])

    # Tier 3: Database & Security Layer
    add_card(slide, Inches(0.8) + (col_w + gap)*2, top_y, col_w, h_y)
    tb3 = slide.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, top_y + Inches(0.2), col_w - Inches(0.4), h_y - Inches(0.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Database & Security Layer"
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = BRAND_BLUE
    
    items3 = [
        ("PostgreSQL RDBMS", "Enterprise relational database with ACID compliance and robust transactional safety."),
        ("20+ Normalized Tables", "Clean 3NF schema eliminating data redundancy across employees, FSUs, and projects."),
        ("Stateless JWT Auth", "JSON Web Tokens with cryptographic signature for secure, scalable session verification."),
        ("Bcrypt Password Hashing", "One-way salted password encryption protecting user credentials against brute-force."),
        ("Audit Trail Logging", "Full timestamped logging of sensitive modifications for compliance.")
    ]
    add_feature_bullets(tf3, items3)
    add_tech_tags(slide, Inches(1.0) + (col_w + gap)*2, top_y + h_y - Inches(0.55), ["PostgreSQL", "JWT", "Bcrypt", "RBAC"])

def build_slide_4_rbac(prs):
    """Slide 4: Role-Based Access Control (RBAC) & Personas"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    add_header(slide, "3. Security & Access Control", "Four-Tier Role-Based Access Control (RBAC)", "Adhering to the NIST RBAC standard to enforce strict data separation and least-privilege principles.")
    add_footer(slide, 4)
    
    grid_w = Inches(5.72)
    grid_h = Inches(2.32)
    gap_x = Inches(0.3)
    gap_y = Inches(0.26)
    x1 = Inches(0.8)
    x2 = x1 + grid_w + gap_x
    y1 = Inches(1.75)
    y2 = y1 + grid_h + gap_y

    roles = [
        ("System Administrator (Admin)", BRAND_BLUE, x1, y1, [
            ("Master Configuration", "Full privilege over system settings, database parameters, and master lookup tables."),
            ("User Provisioning", "Creating employee accounts, resetting passwords, and assigning role credentials."),
            ("Security & Audit Logs", "Monitoring real-time audit logs of critical database changes and access attempts.")
        ]),
        ("Human Resource Manager (HR)", BRAND_ORANGE, x2, y1, [
            ("Workforce Lifecycle", "Onboarding new hires, updating employee records, managing job titles and salary grades."),
            ("FSU Structure Management", "Configuring organizational departments, managers, and headcount targets."),
            ("Payroll & Recruitment", "Calculating monthly department payrolls, approving payslips, and managing candidate ATS pipeline.")
        ]),
        ("Project Manager (PM)", BRAND_BLUE, x1, y2, [
            ("Project Portfolio Control", "Creating projects, tracking delivery milestones, budget health, and completion status."),
            ("Resource Allocation Matrix", "Assigning developers to projects by workload percentage and monitoring team capacity."),
            ("Skill Matching & Leave Review", "Reviewing project team leave requests and scouting engineers via Skill Matrix.")
        ]),
        ("Employee / Developer (Staff)", SUCCESS_GREEN, x2, y2, [
            ("Personal Self-Service", "Managing profile information, updating emergency contacts, and viewing assigned assets."),
            ("Digital Leave Filing", "Submitting leave requests with automatic balance calculation and live approval tracking."),
            ("Confidential Payslip & KPI", "Securely viewing monthly payslips and participating in semi-annual KPI self-appraisals.")
        ])
    ]

    for title, color, rx, ry, items in roles:
        add_card(slide, rx, ry, grid_w, grid_h)
        tb = slide.shapes.add_textbox(rx + Inches(0.2), ry + Inches(0.12), grid_w - Inches(0.4), grid_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        add_feature_bullets(tf, items)

def build_split_module_slide(prs, slide_num, category_tag, title, subtitle, items, tags, screenshot_file, caption):
    """Helper to generate a split-screen module showcase slide (Left Info Card + Right UI Mockup)."""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    add_header(slide, category_tag, title, subtitle)
    add_footer(slide, slide_num)
    
    left_w = Inches(4.8)
    right_w = Inches(6.65)
    gap = Inches(0.28)
    top_y = Inches(1.75)
    h_y = Inches(4.9)

    # Left Feature Card
    add_card(slide, Inches(0.8), top_y, left_w, h_y)
    tb = slide.shapes.add_textbox(Inches(1.05), top_y + Inches(0.2), left_w - Inches(0.5), h_y - Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_head = tf.paragraphs[0]
    p_head.text = "Key Capabilities & Engineering Specs"
    p_head.font.size = Pt(13.5)
    p_head.font.bold = True
    p_head.font.color.rgb = BRAND_BLUE
    
    add_feature_bullets(tf, items)
    if tags:
        add_tech_tags(slide, Inches(1.05), top_y + h_y - Inches(0.52), tags)

    # Right Screenshot Card
    img_path = get_screenshot(screenshot_file)
    add_screenshot_card(slide, img_path, Inches(0.8) + left_w + gap, top_y, right_w, h_y, caption=caption)

def build_slide_5_auth(prs):
    """Slide 5: Module 1 — Authentication Gateway"""
    items = [
        ("Bcrypt One-Way Hashing", "Mật khẩu được mã hóa an toàn bằng thuật toán Bcrypt với Salt ngẫu nhiên, chống tấn công brute-force."),
        ("Quick Role Switcher", "Giao diện tích hợp nút chọn nhanh 4 vai trò (Admin, HR, PM, Employee), hỗ trợ tối đa việc nghiệm thu và demo."),
        ("Stateless JWT Token", "Xác thực phiên qua JWT token lưu an toàn tại Client, đính kèm vào mọi yêu cầu API backend."),
        ("Client-Side Route Guard", "Tự động kiểm tra quyền hạn trước khi render trang; tự động chuyển hướng nếu người dùng truy cập trái phép.")
    ]
    # English translation for slide content:
    items_en = [
        ("Bcrypt Salted Encryption", "Passwords hashed using one-way Bcrypt algorithm with salt factor, guaranteeing credential safety."),
        ("Quick Multi-Role Switcher", "Built-in interactive persona selector allowing instant switching between Admin, HR, PM, and Employee."),
        ("Stateless JWT Bearer Token", "Secure token stored client-side and attached to all API requests for stateless session verification."),
        ("Automated Route Guard", "Client-side navigation interceptor preventing unauthorized access to restricted management views.")
    ]
    build_split_module_slide(
        prs, 5, "Core Module 1 • Gateway & Security",
        "Secure Authentication & Multi-Role Gateway",
        "Dual-layer defense combining Bcrypt password encryption and JWT token-based session management.",
        items_en, ["Bcrypt", "JWT", "Route Guard", "LocalStorage"],
        "01_login.png", "Figure: Multi-Role Secure Login Gateway with Interactive Persona Switcher"
    )

def build_slide_6_dashboard(prs):
    """Slide 6: Module 2 — Executive Dashboard"""
    items_en = [
        ("Real-Time KPI Counters", "Instant executive telemetry displaying active workforce, running projects, pending leaves, and payroll budget."),
        ("FSU Headcount Doughnut", "Dynamic Chart.js visualization mapping employee headcount distribution across Strategic Software Units."),
        ("Delivery Milestone Chart", "Visual bar chart tracking project completion percentages and schedule health across all business units."),
        ("Live Activity Stream", "Real-time feed notifying leadership of recent leave submissions, new recruitments, and allocation updates.")
    ]
    build_split_module_slide(
        prs, 6, "Core Module 2 • Operations Telemetry",
        "Executive Management Dashboard & Analytics",
        "Centralized command portal providing company leadership with live workforce telemetry and project health.",
        items_en, ["Chart.js", "Live Telemetry", "Executive KPI", "Data Visuals"],
        "02_dashboard.png", "Figure: Executive Management Dashboard with Interactive Charts and KPI Counters"
    )

def build_slide_7_departments(prs):
    """Slide 7: Module 3 — Organizational Hierarchy (FSUs)"""
    items_en = [
        ("FSU Specialization Units", "Accurately models FPT Software's Strategic Software Units: GAT (Automotive), GST (IoT/Smart Factory), DBS (Banking), and GAM."),
        ("Unit Leadership Tracking", "Displays appointed FSU Manager, total assigned developer headcount, active projects, and technical focus."),
        ("Hierarchical Data Schema", "Maintains clean relational parent-child mapping across business units and staff rosters."),
        ("Agile Capacity Scaling", "Enables corporate leadership to analyze resource distribution and plan strategic unit expansions.")
    ]
    build_split_module_slide(
        prs, 7, "Core Module 3 • Enterprise Structure",
        "Strategic Software Units (FSU) Architecture",
        "Digital modeling of FPT Software's specialized software delivery units and departmental hierarchy.",
        items_en, ["FSU Modeling", "Org Hierarchy", "Relational Mapping", "Headcount Tracking"],
        "03_departments.png", "Figure: Strategic Software Units (FSU) Organization Management Interface"
    )

def build_slide_8_employees(prs):
    """Slide 8: Module 4 — Employee Lifecycle Management"""
    items_en = [
        ("Comprehensive Directory", "Unified registry of all enterprise staff with search by name, employee code, FSU, and job position."),
        ("Modal-Based CRUD Operations", "Interactive popups for adding new engineers and editing profiles without disruptive page refreshes."),
        ("Full Employment History", "Tracks employment type (Full-time, Intern, Probation), joining dates, contracts, and assigned managers."),
        ("Corporate Credential Mapping", "Directly binds employee records to system authentication accounts and RBAC access roles.")
    ]
    build_split_module_slide(
        prs, 8, "Core Module 4 • Workforce Directory",
        "Employee Lifecycle & Profile Management",
        "Centralized personnel registry supporting multi-criteria search and modal-based CRUD operations.",
        items_en, ["Modal CRUD", "Employee Registry", "Live Search", "RESTful APIs"],
        "04_employees.png", "Figure: Centralized Employee Directory with Multi-Criteria Filtering and Modal Edit"
    )

def build_slide_9_projects(prs):
    """Slide 9: Module 5 — Project Portfolio Management"""
    items_en = [
        ("Portfolio Status Overview", "Categorizes enterprise client projects into Planning, In Progress, On Hold, and Completed."),
        ("Visual Milestone Tracking", "Color-coded progress bars displaying percentage completion and remaining delivery timelines."),
        ("Budget & Delivery Oversight", "Monitors contracted project budgets, planned vs. actual start dates, and designated Project Managers."),
        ("Resource Demand Linking", "Direct bridge between project delivery commitments and developer staffing demands.")
    ]
    build_split_module_slide(
        prs, 9, "Core Module 5 • Delivery Tracking",
        "Project Portfolio Tracking & Milestone Control",
        "Connecting client delivery schedules with technical team demands across all business units.",
        items_en, ["Progress Bar", "Milestone Control", "Project Portfolio", "Delivery Tracking"],
        "06_projects.png", "Figure: Enterprise Project Portfolio Tracking with Milestone Completion Indicators"
    )

def build_slide_10_resources(prs):
    """Slide 10: Module 6 — Dynamic Resource Allocation (Key Innovation)"""
    items_en = [
        ("Interactive Staffing Matrix", "Project Managers assign engineers to projects with specific % workload commitments (e.g. 50% Proj A, 50% Proj B)."),
        ("Burnout Prevention Guard", "Intelligent threshold guard triggers real-time warning if total developer allocation reaches 100%, preventing burnout."),
        ("Bench Cost Minimization", "Instant discovery of unassigned (bench) or under-utilized engineers, accelerating project kickoff."),
        ("Core Business Value", "Directly resolves the critical resource allocation bottleneck identified at FPT Software Da Nang.")
    ]
    build_split_module_slide(
        prs, 10, "Core Module 6 • Key Capstone Innovation",
        "Dynamic Resource Allocation & Bench Optimization",
        "Algorithmic capacity management solving developer over-allocation and idle bench corporate overhead.",
        items_en, ["Resource Matrix", "Burnout Guard", "Bench Discovery", "Capacity Allocation"],
        "07_resources.png", "Figure: Dynamic Resource Allocation Matrix with Workload Percentage Throttling"
    )

def build_slide_11_leave(prs):
    """Slide 11: Module 7 — Paperless Leave Management"""
    items_en = [
        ("100% Digital Workflow", "Replaces physical forms and fragmented email chains with an auditable online leave management engine."),
        ("Smart Working-Day Engine", "Automatically calculates exact working days requested, excluding weekends and official public holidays."),
        ("Two-Tier Approval Chain", "Requests routed seamlessly to Project Manager and HR Manager with real-time status badges."),
        ("Automatic Quota Deduction", "Instantly updates remaining annual leave balance upon final sign-off, maintaining complete audit history.")
    ]
    build_split_module_slide(
        prs, 11, "Core Module 7 • Workflow Automation",
        "Paperless Leave Management & Approval Engine",
        "Streamlined end-to-end digital leave workflow replacing manual paperwork and administrative delays.",
        items_en, ["Leave Workflow", "Two-Tier Approval", "Working-Day Calc", "Quota Tracking"],
        "08_leave.png", "Figure: Digital Leave Request Submission and Multi-Tier Approval Management Portal"
    )

def build_slide_12_payroll(prs):
    """Slide 12: Module 8 — Automated Digital Payroll"""
    items_en = [
        ("Formula-Driven Engine", "Computes gross pay, actual working days, overtime coefficients, allowances, and statutory insurance deductions."),
        ("Batch Department Payroll", "HR administrators can calculate and review payroll for entire software units with a single action."),
        ("Confidential Self-Service", "Individual developers securely view and inspect detailed payslips directly from their personal account."),
        ("Financial Accuracy", "Eliminates human calculation errors, building employee trust through transparent salary breakdowns.")
    ]
    build_split_module_slide(
        prs, 12, "Core Module 8 • Compensation & Benefits",
        "Automated Payroll Engine & Digital Payslips",
        "Accurate financial calculations integrating base salary, attendance, allowances, and statutory taxes.",
        items_en, ["Payroll Engine", "Batch Calculation", "Digital Payslip", "Insurance & Tax"],
        "09_payroll.png", "Figure: Enterprise Payroll Processing and Personal Digital Payslip Interface"
    )

def build_slide_13_performance(prs):
    """Slide 13: Module 9 — KPI Evaluation & Skill Matrix"""
    items_en = [
        ("Dual-Perspective Appraisal", "Combines employee self-assessment with direct manager evaluation on a standardized 100-point rubric."),
        ("Automated Grade Tiering", "Computes official performance rank: Grade A (Top Performer), Grade B, Grade C, or Grade D (Improvement)."),
        ("Competency Skill Matrix", "Maps technical proficiencies across Frontend, Backend, Cloud, DevOps, and QA (Levels 1 to 5)."),
        ("Strategic Squad Assembly", "Enables PMs to scout engineers matching specific technology stacks needed for upcoming client bids.")
    ]
    build_split_module_slide(
        prs, 13, "Core Module 9 • Talent Development",
        "Performance Appraisal (KPI) & Technical Skill Matrix",
        "Objective workforce evaluation coupled with tech stack competency mapping for agile staffing.",
        items_en, ["KPI Rubric", "Performance Grade", "Skill Matrix", "Competency Radar"],
        "11_skills.png", "Figure: Technical Skill Matrix and Competency Proficiency Mapping Interface"
    )

def build_slide_14_auxiliary(prs):
    """Slide 14: Module 10 — Recruitment, Assets & Training"""
    items_en = [
        ("Applicant Tracking (ATS)", "Manages hiring pipeline, resume reviews, interview stages, and direct candidate conversion into employee records."),
        ("Enterprise Asset Custody", "Tracks hardware custody (laptops, test devices, peripherals) assigned to engineers with condition auditing."),
        ("Internal Training Academy", "Maintains technical certification catalog (AWS, Azure, ISTQB, Scrum), enrollment tracking, and completion scoring."),
        ("Holistic ERP Experience", "Provides an integrated enterprise ecosystem covering every aspect of operational administration.")
    ]
    build_split_module_slide(
        prs, 14, "Core Module 10 • Enterprise Operations",
        "Recruitment ATS, Asset Tracking & Internal Academy",
        "Supporting enterprise modules delivering an all-in-one operational ecosystem for software centers.",
        items_en, ["Recruitment ATS", "Asset Custody", "Training Academy", "Enterprise ERP"],
        "13_recruitment.png", "Figure: Candidate Recruitment Pipeline and Applicant Tracking System (ATS)"
    )

def build_slide_15_impact(prs):
    """Slide 15: Key Technical Achievements & Business Value Delivered"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    add_header(slide, "4. Results & Business Impact", "Measurable Business Impact & Technical Highlights", "Transforming operational efficiency while establishing a rock-solid, production-grade technical foundation.")
    add_footer(slide, 15)
    
    card_w = Inches(5.72)
    card_h = Inches(2.32)
    gap_x = Inches(0.3)
    gap_y = Inches(0.26)
    x1 = Inches(0.8)
    x2 = x1 + card_w + gap_x
    y1 = Inches(1.75)
    y2 = y1 + card_h + gap_y

    metrics = [
        ("100% Paperless Digitalization", BRAND_BLUE, x1, y1, [
            ("Zero Physical Forms", "Completely eliminated manual paperwork and disconnected spreadsheets across departments."),
            ("Single-Pane Portal", "Unified all HR, project staffing, payroll, and asset workflows onto one consolidated platform."),
            ("Seamless Collaboration", "Accelerated communication between PMs, HR executives, and engineering team members.")
        ]),
        ("Optimal Staffing & Bench Control", BRAND_ORANGE, x2, y1, [
            ("Zero Developer Burnout", "Automated 100% capacity threshold guard successfully blocks dangerous over-allocation."),
            ("Rapid Bench Discovery", "Project Managers locate available engineers with required skills in seconds instead of days."),
            ("Cost Optimization", "Directly reduces corporate idle bench overhead through transparent workforce allocation.")
        ]),
        ("Sub-Second System Performance", BRAND_BLUE, x1, y2, [
            ("Lightweight Client Footprint", "Built with vanilla ES6+ and modern CSS, eliminating heavy client-side bundle load times."),
            ("High-Throughput Node.js", "Asynchronous non-blocking backend architecture serving REST API requests with sub-300ms latency."),
            ("Optimized Relational Schema", "PostgreSQL database with indexed foreign keys and normalized 3NF relational tables.")
        ]),
        ("Enterprise-Grade Security", SUCCESS_GREEN, x2, y2, [
            ("NIST RBAC Standards", "Granular four-tier role access model safeguarding confidential salary and personnel data."),
            ("Bcrypt Salted Encryption", "State-of-the-art password protection preventing offline dictionary and rainbow-table attacks."),
            ("Cryptographic JWT Tokens", "Stateless session validation with automated route guards and security audit logging.")
        ])
    ]

    for title, color, rx, ry, items in metrics:
        add_card(slide, rx, ry, card_w, card_h)
        tb = slide.shapes.add_textbox(rx + Inches(0.2), ry + Inches(0.12), card_w - Inches(0.4), card_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color
        add_feature_bullets(tf, items)

def build_slide_16_conclusion(prs):
    """Slide 16: Conclusion, Lessons Learned & Future Roadmap (Dark Theme)"""
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    
    # Dark Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DARK
    bg.line.fill.background()
    
    add_header(slide, "5. Conclusion & Roadmap", "Project Summary, Lessons Learned & Future Vision", "Reflecting on real-world engineering experiences at FPT Software Da Nang and future technological evolution.", is_dark=True)
    add_footer(slide, 16, is_dark=True)
    
    card_w = Inches(5.72)
    card_h = Inches(3.4)
    top_y = Inches(1.75)
    gap = Inches(0.3)

    # Left Card: Lessons Learned
    add_card(slide, Inches(0.8), top_y, card_w, card_h, bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb1 = slide.shapes.add_textbox(Inches(1.05), top_y + Inches(0.2), card_w - Inches(0.5), card_h - Inches(0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Key Internship Takeaways"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    
    lessons = [
        ("CMMI Level 5 Software Processes", "Acquired practical understanding of standard enterprise software development lifecycles and quality gates."),
        ("Fullstack Architecture Mastery", "Gained deep competency designing RESTful APIs, relational schema normalization (3NF), and responsive UI layouts."),
        ("Security-First Engineering", "Implemented enterprise-grade protection: stateless authentication, route guards, and audit trail logging."),
        ("Enterprise Problem Solving", "Translated real corporate pain points at FPT Complex into concrete, deployable software features.")
    ]
    add_feature_bullets(tf1, lessons, is_dark=True)

    # Right Card: Future Roadmap
    add_card(slide, Inches(0.8) + card_w + gap, top_y, card_w, card_h, bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb2 = slide.shapes.add_textbox(Inches(1.05) + card_w + gap, top_y + Inches(0.2), card_w - Inches(0.5), card_h - Inches(0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Future Technological Roadmap"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = BRAND_ORANGE
    
    roadmap = [
        ("Cross-Platform Mobile App", "Developing native mobile companion apps (Flutter / React Native) for instant push notifications and on-the-go approvals."),
        ("AI-Powered Resource Matcher", "Integrating LLM-based CV parsing to automatically match candidate skill profiles with incoming client project bids."),
        ("Real-Time WebSockets", "Deploying WebSocket channels for live in-app collaborative staffing and instant leave ticket status updates."),
        ("Automated CI/CD Deployment", "Containerizing the backend with Docker and establishing automated GitHub Actions testing and deployment pipelines.")
    ]
    add_feature_bullets(tf2, roadmap, is_dark=True)

    # Bottom Acknowledgments & Q&A Banner
    banner = add_card(slide, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.35), bg_color=NAVY_CARD, border_color=RGBColor(30, 58, 110))
    tb3 = slide.shapes.add_textbox(Inches(1.1), Inches(5.45), Inches(11.133), Inches(1.15))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    
    p3 = tf3.paragraphs[0]
    p3.text = "THANK YOU FOR YOUR ATTENTION!"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = WHITE
    
    p4 = tf3.add_paragraph()
    p4.text = "Sincere gratitude to Vietnam - Korea University (VKU), M.Sc. Nguyen Ngoc Huyen Tran, and FPT Software Da Nang mentors for guidance and support."
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(10.5)
    p4.font.color.rgb = RGBColor(203, 213, 225)
    p4.space_before = Pt(4)
    
    p5 = tf3.add_paragraph()
    p5.text = "Q & A  •  OPEN FOR QUESTIONS & VALUABLE FEEDBACK"
    p5.alignment = PP_ALIGN.CENTER
    p5.font.size = Pt(12)
    p5.font.bold = True
    p5.font.color.rgb = BRAND_ORANGE
    p5.space_before = Pt(4)

# ==============================================================================
# MAIN GENERATOR
# ==============================================================================

def main():
    print(f"Creating PowerPoint Presentation: {PPTX_PATH}")
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    print("Building Slide 1: Title Slide (Dark Theme)...")
    build_slide_1_title(prs)
    
    print("Building Slide 2: Motivation & Problem Statement...")
    build_slide_2_motivation(prs)
    
    print("Building Slide 3: System Architecture & Tech Stack...")
    build_slide_3_architecture(prs)
    
    print("Building Slide 4: Role-Based Access Control (RBAC)...")
    build_slide_4_rbac(prs)
    
    print("Building Slide 5: Core Module 1 — Authentication Gateway...")
    build_slide_5_auth(prs)
    
    print("Building Slide 6: Core Module 2 — Executive Dashboard...")
    build_slide_6_dashboard(prs)
    
    print("Building Slide 7: Core Module 3 — Organizational Hierarchy (FSUs)...")
    build_slide_7_departments(prs)
    
    print("Building Slide 8: Core Module 4 — Employee Lifecycle...")
    build_slide_8_employees(prs)
    
    print("Building Slide 9: Core Module 5 — Project Portfolio...")
    build_slide_9_projects(prs)
    
    print("Building Slide 10: Core Module 6 — Dynamic Resource Allocation...")
    build_slide_10_resources(prs)
    
    print("Building Slide 11: Core Module 7 — Paperless Leave Management...")
    build_slide_11_leave(prs)
    
    print("Building Slide 12: Core Module 8 — Automated Digital Payroll...")
    build_slide_12_payroll(prs)
    
    print("Building Slide 13: Core Module 9 — KPI Performance & Skill Matrix...")
    build_slide_13_performance(prs)
    
    print("Building Slide 14: Core Module 10 — Recruitment, Assets & Training...")
    build_slide_14_auxiliary(prs)
    
    print("Building Slide 15: Key Technical Achievements & Business Impact...")
    build_slide_15_impact(prs)
    
    print("Building Slide 16: Conclusion, Lessons Learned & Roadmap (Dark Theme)...")
    build_slide_16_conclusion(prs)
    
    # Save presentation
    prs.save(PPTX_PATH)
    print(f"\n=======================================================")
    print(f"SUCCESSFULLY GENERATED 16 SLIDES TO: {PPTX_PATH}")
    print(f"=======================================================")

if __name__ == "__main__":
    main()
