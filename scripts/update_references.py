# -*- coding: utf-8 -*-
"""
Script: update_references.py
Author: Antigravity
Purpose:
  1. Restore clean document from backup to ensure reproducible clean run.
  2. Add 20 scholarly in-text citations ([1] to [6]) formatted as superscript square brackets 
     at appropriate sentences/phrases across the report body.
  3. Strictly avoid inserting citations into headings, captions, or Table of Contents.
  4. Populate the 'TÀI LIỆU THAM KHẢO' section with the 6 selected references in IEEE format:
     [1] Sommerville (Software Engineering)
     [2] Ferraiolo et al. (Role-Based Access Control - RBAC NIST Standard)
     [3] Fielding (RESTful API Architecture)
     [4] Jones et al. (JSON Web Token - RFC 7519)
     [5] Provos & Mazières (Bcrypt Password Hashing)
     [6] Tilkov & Vinoski (Node.js Network Programs)
  5. Automate Microsoft Word via COM to update TOC, fields, and save cleanly.
"""

import os
import sys
import shutil
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client

DOC_PATH = r"D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\2. Mau bao cao thuc tap.docx"
BAK_PATH = r"D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\2. Mau bao cao thuc tap.docx.pre_references.bak"

def insert_superscript_citation(p, target_phrase, citation_str):
    """
    Inserts a superscript citation run immediately after target_phrase in paragraph p.
    Handles non-breaking spaces (\xa0) transparently.
    Preserves existing run formatting.
    """
    norm_full = p.text.replace('\xa0', ' ')
    norm_target = target_phrase.replace('\xa0', ' ')
    
    if norm_target not in norm_full:
        return False
    if citation_str in p.text:
        return False  # Already cited in this paragraph
        
    start_pos = norm_full.find(norm_target)
    end_pos = start_pos + len(norm_target)
    
    # Locate which run contains end_pos
    curr_pos = 0
    target_run = None
    split_offset = -1
    
    for r in p.runs:
        r_len = len(r.text)
        if curr_pos <= end_pos <= curr_pos + r_len:
            target_run = r
            split_offset = end_pos - curr_pos
            break
        curr_pos += r_len
        
    if not target_run:
        return False
        
    # Superscript run XML (Times New Roman, sz=26 -> 13pt base, superscript vertical alignment)
    sup_xml = f'''
        <w:r {nsdecls("w")}>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:vertAlign w:val="superscript"/>
                <w:sz w:val="26"/>
                <w:szCs w:val="26"/>
            </w:rPr>
            <w:t>{citation_str}</w:t>
        </w:r>
    '''
    r_sup_elem = parse_xml(sup_xml)
    
    if split_offset == len(target_run.text):
        target_run._element.addnext(r_sup_elem)
    else:
        text_before = target_run.text[:split_offset]
        text_after = target_run.text[split_offset:]
        target_run.text = text_before
        
        # Clone run for text_after to preserve original font properties
        after_elem = parse_xml(target_run._element.xml)
        for t in after_elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
            t.text = text_after
            break
            
        target_run._element.addnext(r_sup_elem)
        r_sup_elem.addnext(after_elem)
        
    return True

def create_reference_paragraph_xml(label, authors, title_or_article, is_article_quotes, source_title, is_source_italic, details):
    """
    Creates an IEEE formatted reference paragraph XML element with hanging indentation.
    """
    ns = nsdecls("w")
    runs_xml = []
    
    # 1. Label run e.g. [1]
    runs_xml.append(f'''
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:sz w:val="26"/>
                <w:szCs w:val="26"/>
            </w:rPr>
            <w:t xml:space="preserve">{label} </w:t>
        </w:r>
    ''')
    
    # 2. Authors
    if authors:
        runs_xml.append(f'''
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:sz w:val="26"/>
                    <w:szCs w:val="26"/>
                </w:rPr>
                <w:t xml:space="preserve">{authors} </w:t>
            </w:r>
        ''')
        
    # 3. Title / Article
    if title_or_article:
        rpr = '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="26"/><w:szCs w:val="26"/>'
        if not is_article_quotes:
            rpr += '<w:i/><w:iCs/>'
            t_text = f"{title_or_article},"
        else:
            t_text = f'"{title_or_article},"'
        runs_xml.append(f'''
            <w:r>
                <w:rPr>{rpr}</w:rPr>
                <w:t xml:space="preserve">{t_text} </w:t>
            </w:r>
        ''')
        
    # 4. Source / Journal / Conference / Dissertation
    if source_title:
        rpr = '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="26"/><w:szCs w:val="26"/>'
        if is_source_italic:
            rpr += '<w:i/><w:iCs/>'
        runs_xml.append(f'''
            <w:r>
                <w:rPr>{rpr}</w:rPr>
                <w:t xml:space="preserve">{source_title} </w:t>
            </w:r>
        ''')
        
    # 5. Details (edition, volume, pages, year)
    if details:
        runs_xml.append(f'''
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:sz w:val="26"/>
                    <w:szCs w:val="26"/>
                </w:rPr>
                <w:t xml:space="preserve">{details}</w:t>
            </w:r>
        ''')

    runs_content = "".join(runs_xml)
    p_xml = f'''
    <w:p {ns}>
        <w:pPr>
            <w:pStyle w:val="Normal"/>
            <w:ind w:left="425" w:hanging="425"/>
            <w:jc w:val="both"/>
            <w:spacing w:line="280" w:lineRule="auto" w:after="100"/>
        </w:pPr>
        {runs_content}
    </w:p>
    '''
    return parse_xml(p_xml)

def main():
    # 0. Ensure clean state from backup
    if os.path.exists(BAK_PATH):
        print(f"Restoring clean baseline from: {BAK_PATH}")
        shutil.copy2(BAK_PATH, DOC_PATH)
    else:
        print("Warning: Backup file not found, proceeding with current document.")

    print(f"Loading document: {DOC_PATH}")
    doc = docx.Document(DOC_PATH)

    # 1. Mapping of in-text citations to phrases (6 references)
    citations_to_apply = [
        # [1] Sommerville (Software Engineering)
        ("chuẩn hóa toàn bộ vòng đời quản trị nhân sự trong tổ chức", "[1]"),
        ("thiết kế cơ sở dữ liệu tối ưu.", "[1]"),
        ("kiến thức nền tảng về công nghệ phần mềm", "[1]"),

        # [2] Ferraiolo et al. (Role-Based Access Control - RBAC NIST Standard)
        ("kiểm soát truy cập dựa trên vai trò (RBAC)", "[2]"),
        ("kiểm soát truy cập (RBAC):", "[2]"),
        ("chuyển hướng chính xác đến khu vực làm việc tương ứng.", "[2]"),
        ("kiểm soát truy cập nghiêm ngặt theo bốn vai trò", "[2]"),

        # [3] Fielding (RESTful API Architecture)
        ("(Routes, Controllers, Models, Middlewares)", "[3]"),
        ("tiêu chuẩn kiến trúc RESTful API", "[3]"),

        # [4] Jones et al. (JSON Web Token - RFC 7519)
        ("xác thực bằng mã thông báo JWT", "[4]"),
        ("cơ chế mã thông báo JSON Web Token (JWT)", "[4]"),
        ("định danh phiên bằng JSON Web Token", "[4]"),
        ("xác thực phiên làm việc bằng JSON Web Token", "[4]"),

        # [5] Provos & Mazières (Bcrypt Password Hashing)
        ("thông qua mã hóa mật khẩu Bcrypt,", "[5]"),
        ("mã hóa một chiều bằng thuật toán an toàn Bcrypt.", "[5]"),
        ("mã hóa mật khẩu Bcrypt và", "[5]"),
        ("cơ chế mã hóa mật khẩu Bcrypt,", "[5]"),

        # [6] Tilkov & Vinoski (Node.js Network Programs)
        ("phát triển bằng môi trường Node.js", "[6]"),
        ("xây dựng trên nền tảng Node.js", "[6]"),
        ("dựa trên nền tảng Node.js,", "[6]"),
    ]

    applied_count = 0
    # Apply citations ONLY to body paragraphs (index >= 174 and < 386)
    # Strictly exclude Headings, Titles, Subtitles, and Captions
    for target_phrase, citation_str in citations_to_apply:
        for idx in range(174, min(386, len(doc.paragraphs))):
            p = doc.paragraphs[idx]
            if p.style.name.startswith("Heading") or p.style.name in ["Title", "Subtitle", "Caption"]:
                continue
            if insert_superscript_citation(p, target_phrase, citation_str):
                applied_count += 1
                print(f"Applied citation {citation_str:5s} at P{idx:3d} after: '{target_phrase[:45]}...'")

    print(f"\n===> Total in-text citations successfully applied: {applied_count} / {len(citations_to_apply)}")

    # 2. Populate 'TÀI LIỆU THAM KHẢO' Section
    ref_heading_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "TÀI LIỆU THAM KHẢO":
            ref_heading_idx = i
            break

    if ref_heading_idx is None:
        print("Error: Heading 'TÀI LIỆU THAM KHẢO' not found!")
        return

    print(f"Found 'TÀI LIỆU THAM KHẢO' at paragraph {ref_heading_idx}")

    body = doc._element.body
    ref_heading_elem = doc.paragraphs[ref_heading_idx]._element
    
    # Locate the final note paragraph
    final_note_elem = None
    for i in range(ref_heading_idx + 1, len(doc.paragraphs)):
        if "Lưu ý rằng:" in doc.paragraphs[i].text:
            final_note_elem = doc.paragraphs[i]._element
            break

    # Remove placeholder empty paragraphs
    curr = ref_heading_elem.getnext()
    elems_to_remove = []
    while curr is not None and curr != final_note_elem:
        elems_to_remove.append(curr)
        curr = curr.getnext()

    for el in elems_to_remove:
        body.remove(el)
    print(f"Removed {len(elems_to_remove)} placeholder empty paragraphs in reference section.")

    # 6 References Data matching IEEE standard and uploaded user image
    references_data = [
        {
            "label": "[1]",
            "authors": "I. Sommerville,",
            "title_or_article": "Software Engineering",
            "is_article_quotes": False,
            "source_title": "",
            "is_source_italic": False,
            "details": "10th ed., Boston, MA: Pearson Education, 2015."
        },
        {
            "label": "[2]",
            "authors": "D. F. Ferraiolo, R. Sandhu, S. Gavrila, D. R. Kuhn, and R. Chandramouli,",
            "title_or_article": "Proposed NIST standard for role-based access control",
            "is_article_quotes": True,
            "source_title": "ACM Transactions on Information and System Security,",
            "is_source_italic": True,
            "details": "vol. 4, no. 3, pp. 224–274, 2001."
        },
        {
            "label": "[3]",
            "authors": "R. T. Fielding,",
            "title_or_article": "Architectural Styles and the Design of Network-based Software Architectures",
            "is_article_quotes": True,
            "source_title": "Ph.D. dissertation,",
            "is_source_italic": False,
            "details": "Department of Information and Computer Science, University of California, Irvine, CA, 2000."
        },
        {
            "label": "[4]",
            "authors": "M. Jones, J. Bradley, and N. Sakimura,",
            "title_or_article": "JSON Web Token (JWT)",
            "is_article_quotes": True,
            "source_title": "RFC 7519,",
            "is_source_italic": True,
            "details": "Internet Engineering Task Force (IETF), pp. 1–30, 2015."
        },
        {
            "label": "[5]",
            "authors": "N. Provos and D. Mazières,",
            "title_or_article": "A future-adaptable password scheme",
            "is_article_quotes": True,
            "source_title": "in Proceedings of the FREENIX Track: 1999 USENIX Annual Technical Conference,",
            "is_source_italic": True,
            "details": "Monterey, CA, 1999, pp. 81–91."
        },
        {
            "label": "[6]",
            "authors": "S. Tilkov and S. Vinoski,",
            "title_or_article": "Node.js: Using JavaScript to build high-performance network programs",
            "is_article_quotes": True,
            "source_title": "IEEE Internet Computing,",
            "is_source_italic": True,
            "details": "vol. 14, no. 6, pp. 80–83, 2010."
        }
    ]

    # Insert references right before final_note_elem (or at end of doc)
    target_insert_elem = final_note_elem if final_note_elem is not None else None
    
    for ref in references_data:
        p_ref = create_reference_paragraph_xml(
            label=ref["label"],
            authors=ref["authors"],
            title_or_article=ref["title_or_article"],
            is_article_quotes=ref["is_article_quotes"],
            source_title=ref["source_title"],
            is_source_italic=ref["is_source_italic"],
            details=ref["details"]
        )
        if target_insert_elem is not None:
            target_insert_elem.addprevious(p_ref)
        else:
            body.append(p_ref)

    print("Appended 6 IEEE formatted reference entries successfully!")

    # Save document with python-docx
    doc.save(DOC_PATH)
    print("Document saved successfully via python-docx!")

    # 3. Use Word COM to refresh fields and TOC
    print("Opening Microsoft Word COM to update fields and TOC...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(os.path.abspath(DOC_PATH))
        
        # Update Table of Contents
        for i in range(1, wdoc.TablesOfContents.Count + 1):
            wdoc.TablesOfContents(i).Update()
        print("Table of Contents updated!")

        # Update all fields in document
        wdoc.Fields.Update()
        print("All fields updated!")
        
        # Check total pages
        page_count = wdoc.ComputeStatistics(2) # 2 = wdStatisticPages
        print(f"Total document page count: {page_count} pages")
        
        wdoc.Save()
        print("Word document saved successfully via Word COM!")
    finally:
        wdoc.Close(False)
        word.Quit()

    print("\n=======================================================")
    print("ALL REFERENCE TASKS COMPLETED PERFECTLY!")
    print("=======================================================")

if __name__ == "__main__":
    main()
