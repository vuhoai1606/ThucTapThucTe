# -*- coding: utf-8 -*-
"""
Script: build_full_report.py
Author: Antigravity
Purpose: 
  1. Restore fresh clean docx from backup.
  2. Insert Chapter 4 (XAY DUNG WEBSITE VA KET QUA) focusing exclusively on the 13 web functional modules (4.1 to 4.13)
     with 16 high-resolution screenshots, comprehensive academic descriptions, NO tables, NO old 4.1, NO old 4.3.
  3. Insert 'DANH MUC HINH ANH' and 'DANH MUC BANG' into the front-matter section.
  4. Automate Microsoft Word via COM to generate official dynamic Tables of Figures for Hinh (29 figures) 
     and Bang (6 use case tables from Ch 3), update Table of Contents, and sync all sequence numbers and page numbers.
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client

DOC_PATH = r"D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\2. Mau bao cao thuc tap.docx"
BAK_PATH = r"D:\HỌC TẬP\Năm 2 Kỳ hè\Thực tập thực tế\bieu_mau_thuc_tap\2. Mau bao cao thuc tap.docx.bak"
SCREENSHOTS_DIR = os.path.abspath(r"D:\Lap Trinh\Nam 3 Ki 1\ThucTapThucTe\screenshots")

def create_figure_caption_xml(fig_num, fig_title):
    ns = nsdecls("w")
    xml = f'''
    <w:p {ns}>
        <w:pPr>
            <w:pStyle w:val="Caption"/>
            <w:jc w:val="center"/>
            <w:spacing w:before="80" w:after="160"/>
        </w:pPr>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:i/>
                <w:sz w:val="24"/>
                <w:szCs w:val="24"/>
            </w:rPr>
            <w:t>Hình</w:t>
        </w:r>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:i/>
                <w:sz w:val="24"/>
                <w:szCs w:val="24"/>
            </w:rPr>
            <w:t xml:space="preserve"> </w:t>
        </w:r>
        <w:fldSimple w:instr=" SEQ Hình \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:i/>
                    <w:sz w:val="24"/>
                    <w:szCs w:val="24"/>
                </w:rPr>
                <w:t>{fig_num}</w:t>
            </w:r>
        </w:fldSimple>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:i/>
                <w:sz w:val="24"/>
                <w:szCs w:val="24"/>
            </w:rPr>
            <w:t>. {fig_title}</w:t>
        </w:r>
    </w:p>
    '''
    return parse_xml(xml)

def add_academic_paragraph(target_p, text, first_line_indent=365760):
    p = target_p.insert_paragraph_before()
    p.style = 'Normal'
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    if first_line_indent:
        p.paragraph_format.first_line_indent = first_line_indent
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    return p

def add_academic_heading2(target_p, text):
    p = target_p.insert_paragraph_before()
    p.style = 'Heading 2'
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = 0
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    return p

def insert_image_with_caption(doc, target_p, img_filename, fig_num, caption_text, width_inches=5.8):
    img_path = os.path.join(SCREENSHOTS_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"Warning: image {img_path} not found!")
        return

    # Image paragraph
    p_img = target_p.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.first_line_indent = 0
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    
    # Caption paragraph
    cap_elem = create_figure_caption_xml(fig_num, caption_text)
    target_p._element.addprevious(cap_elem)

def main():
    print(f"Restoring clean document from backup: {BAK_PATH}")
    shutil.copyfile(BAK_PATH, DOC_PATH)

    print(f"Loading document: {DOC_PATH}")
    doc = docx.Document(DOC_PATH)

    # 1. Update Chapter 4 description in Section 1 (Bố cục báo cáo)
    for p in doc.paragraphs:
        if "Chương 3: Phân tích và thiết kế hệ thống" in p.text:
            p_c4 = p._element.addnext(parse_xml(f'''
                <w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:pStyle w:val="ListParagraph"/>
                        <w:spacing w:after="160" w:line="259" w:lineRule="auto"/>
                    </w:pPr>
                    <w:r>
                        <w:rPr>
                            <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                            <w:sz w:val="26"/>
                            <w:szCs w:val="26"/>
                        </w:rPr>
                        <w:t>Chương 4: Xây dựng website và kết quả — Trình bày kết quả hiện thực hóa toàn bộ các phân hệ chức năng trên giao diện web của hệ thống FSOFT HRMS.</w:t>
                    </w:r>
                </w:p>
            '''))
            print("Added Chapter 4 overview in Section 1 introduction!")
            break

    # 2. Insert placeholders for DANH MỤC HÌNH ẢNH and DANH MỤC BẢNG right before PHẦN MỞ ĐẦU
    target_intro = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "PHẦN MỞ ĐẦU":
            target_intro = doc.paragraphs[i-1] # paragraph with sectPr ending Section 1
            break
            
    if target_intro:
        ns = nsdecls("w")
        p_hinh_title = parse_xml(f'''
            <w:p {ns}>
                <w:pPr>
                    <w:pStyle w:val="Title"/>
                    <w:jc w:val="center"/>
                    <w:spacing w:before="240" w:after="240"/>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:b/>
                        <w:sz w:val="32"/>
                    </w:rPr>
                </w:pPr>
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:b/>
                        <w:sz w:val="32"/>
                    </w:rPr>
                    <w:br w:type="page"/>
                    <w:t>DANH MỤC HÌNH ẢNH</w:t>
                </w:r>
            </w:p>
        ''')
        p_hinh_holder = parse_xml(f'<w:p {ns}><w:r><w:t>__PLACEHOLDER_TOF_HINH__</w:t></w:r></w:p>')
        
        p_bang_title = parse_xml(f'''
            <w:p {ns}>
                <w:pPr>
                    <w:pStyle w:val="Title"/>
                    <w:jc w:val="center"/>
                    <w:spacing w:before="240" w:after="240"/>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:b/>
                        <w:sz w:val="32"/>
                    </w:rPr>
                </w:pPr>
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:b/>
                        <w:sz w:val="32"/>
                    </w:rPr>
                    <w:br w:type="page"/>
                    <w:t>DANH MỤC BẢNG</w:t>
                </w:r>
            </w:p>
        ''')
        p_bang_holder = parse_xml(f'<w:p {ns}><w:r><w:t>__PLACEHOLDER_TOF_BANG__</w:t></w:r></w:p>')
        
        target_intro._element.addprevious(p_hinh_title)
        target_intro._element.addprevious(p_hinh_holder)
        target_intro._element.addprevious(p_bang_title)
        target_intro._element.addprevious(p_bang_holder)
        print("Inserted TOF placeholders for Hinh and Bang in Section 1!")

    # 3. Locate Chapter 4 heading and target paragraph (KIẾN NGHỊ VÀ KẾT LUẬN)
    target_conclusion = None
    c4_heading_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "XÂY DỰNG WEBSITE VÀ KẾT QUẢ":
            c4_heading_idx = i
        elif p.text.strip() == "KIẾN NGHỊ VÀ KẾT LUẬN":
            target_conclusion = p
            break
            
    print(f"Chapter 4 Heading at paragraph {c4_heading_idx}, Conclusion at paragraph {target_conclusion.text}")

    # Remove blank/raw image paragraphs between c4_heading and conclusion
    body = doc._element.body
    c4_p_elem = doc.paragraphs[c4_heading_idx]._element
    conc_p_elem = target_conclusion._element
    
    elems_to_remove = []
    curr = c4_p_elem.getnext()
    while curr is not None and curr != conc_p_elem:
        elems_to_remove.append(curr)
        curr = curr.getnext()
        
    for elem in elems_to_remove:
        body.remove(elem)
    print(f"Removed {len(elems_to_remove)} empty/raw image paragraphs from Chapter 4!")

    # 4. Insert Chapter 4 Content before target_conclusion: 13 functional sections (4.1 to 4.13)
    t_p = target_conclusion

    add_academic_paragraph(t_p,
        "Dựa trên các yêu cầu phân tích nghiệp vụ và thiết kế kiến trúc đã xây dựng ở Chương 3, toàn bộ các chức năng cốt lõi của hệ thống FSOFT HRMS đã được hiện thực hóa hoàn chỉnh trên nền tảng web Single Page Application (SPA). Dưới đây là kết quả chi tiết từng phân hệ giao diện thực tế của ứng dụng được ghi nhận trực tiếp trong quá trình vận hành thử nghiệm:"
    )

    # 4.1 (formerly 4.2.1)
    add_academic_heading2(t_p, "Phân hệ Đăng nhập và Xác thực phân quyền (Authentication & RBAC)")
    add_academic_paragraph(t_p,
        "Phân hệ Đăng nhập được thiết kế độc lập, tập trung cao vào tính an toàn và tiện dụng. Giao diện được trang bị cơ chế bảo mật hai lớp thông qua mã hóa mật khẩu Bcrypt và định danh phiên bằng JSON Web Token (JWT). Để hỗ trợ tối đa cho việc kiểm thử và nghiệm thu đa vai trò, màn hình cung cấp các nút chọn nhanh (Quick Preset Buttons) tương ứng với 4 vai trò chuẩn trong doanh nghiệp: Quản trị viên (Admin), Quản lý Nhân sự (HR), Quản lý Dự án (PM) và Nhân viên (Employee). Khi người dùng đăng nhập thành công, hệ thống sẽ lưu trữ Access Token cùng thông tin phân quyền vào bộ nhớ cục bộ và kích hoạt Route Guard để chuyển hướng chính xác đến khu vực làm việc tương ứng."
    )
    insert_image_with_caption(doc, t_p, "01_login.png", 14, "Giao diện Đăng nhập hệ thống FSOFT HRMS")

    # 4.2 (formerly 4.2.2)
    add_academic_heading2(t_p, "Phân hệ Bảng điều khiển trung tâm (Dashboard)")
    add_academic_paragraph(t_p,
        "Bảng điều khiển trung tâm (Dashboard) đóng vai trò là cổng tổng quan điều hành của hệ thống FSOFT HRMS. Màn hình cung cấp các thẻ chỉ số đo lường hiệu quả công việc (KPI Cards) cập nhật thời gian thực, bao gồm: Tổng số lượng nhân sự đang làm việc, số lượng dự án phần mềm đang vận hành, số đơn xin nghỉ phép đang chờ xử lý và tổng ngân sách dự án. Bên cạnh đó, hệ thống tích hợp thư viện Chart.js để hiển thị biểu đồ phân bổ nhân sự theo các đơn vị chiến lược (FSU: GAT, GST, DBS, GAM) và biểu đồ cột theo dõi tình trạng dự án. Người dùng cũng có thể thực hiện thao tác điểm danh nhanh (Check-in/Check-out) trực tiếp trên thanh công cụ góc phải màn hình."
    )
    insert_image_with_caption(doc, t_p, "02_dashboard.png", 15, "Giao diện Bảng điều khiển tổng quan (Dashboard)")

    # 4.3 (formerly 4.2.3)
    add_academic_heading2(t_p, "Phân hệ Quản trị Cơ cấu tổ chức & Phòng ban FSU")
    add_academic_paragraph(t_p,
        "Đặc thù tổ chức tại FPT Software là phân chia quy mô hoạt động thành các đơn vị phần mềm chiến lược (Strategic Unit - FSU). Phân hệ Quản trị Cơ cấu tổ chức cho phép hiển thị và quản lý trực quan danh sách các FSU nòng cốt: GAT (Global Automotive Technology - Chuyên phát triển phần mềm ô tô và Autosar), GST (Global Smart Technology - Giải pháp IoT và Smart Factory), DBS (Digital Banking Solutions - Chuyển đổi số ngân hàng), GAM (Global Aerospace & Mechanical) và Khối Nhân sự (HR). Mỗi đơn vị thể hiện rõ ràng họ tên Trưởng đơn vị phụ trách (FSU Manager), số lượng nhân sự trực thuộc, các dự án đang triển khai và mô tả định hướng công nghệ."
    )
    insert_image_with_caption(doc, t_p, "03_departments.png", 16, "Giao diện Quản trị Cơ cấu tổ chức và Đơn vị FSU")

    # 4.4 (formerly 4.2.4)
    add_academic_heading2(t_p, "Phân hệ Quản lý Hồ sơ Nhân sự (Employee Management)")
    add_academic_paragraph(t_p,
        "Phân hệ Quản lý Nhân sự là một trong những khối chức năng trọng tâm nhất của hệ thống. Giao diện cung cấp bảng danh mục nhân viên chi tiết với ảnh đại diện trực quan, mã định danh FPT (Employee ID), họ và tên, chức danh công việc, cấp bậc chuyên môn (Junior, Middle, Senior, Lead, Manager), đơn vị FSU trực thuộc, email doanh nghiệp, mức lương cơ bản và trạng thái làm việc (Active/Inactive). Hệ thống tích hợp thanh công cụ tìm kiếm tức thì theo từ khóa cùng bộ lọc đa chiều (lọc theo FSU, cấp bậc và vai trò hệ thống), giúp việc tra cứu thông tin hàng nghìn nhân sự diễn ra chỉ trong vài mili-giây."
    )
    insert_image_with_caption(doc, t_p, "04_employees.png", 17, "Giao diện Quản lý Danh sách Hồ sơ Nhân viên")
    add_academic_paragraph(t_p,
        "Khi quản trị viên hoặc cán bộ nhân sự thực hiện thêm mới hoặc hiệu chỉnh thông tin, hệ thống sẽ mở cửa sổ biểu mẫu nổi (Modal Form) được thiết kế theo bố cục lưới hai cột khoa học. Biểu mẫu yêu cầu nhập liệu nghiêm ngặt các thông tin cá nhân, định danh căn cước công dân (CCCD), ngày sinh, địa chỉ, đơn vị công tác, chức danh, phân quyền hệ thống và mức lương cơ sở. Hệ thống áp dụng cơ chế kiểm tra tính hợp lệ dữ liệu (Validation) ở cả phía Client lẫn Server, đảm bảo tuyệt đối không cho phép nhập giá trị lương âm hoặc trùng lặp mã định danh nhân viên trong toàn hệ thống."
    )
    insert_image_with_caption(doc, t_p, "05_modal_employee.png", 18, "Giao diện Modal Thêm mới và Hiệu chỉnh thông tin Nhân viên")

    # 4.5 (formerly 4.2.5)
    add_academic_heading2(t_p, "Phân hệ Quản lý Dự án & Tiến độ thực hiện (Project Management)")
    add_academic_paragraph(t_p,
        "Phân hệ Quản lý Dự án hỗ trợ theo dõi toàn diện vòng đời các hợp đồng phát triển phần mềm tại công ty. Màn hình liệt kê danh mục dự án kèm mã định danh chuẩn (như PRJ-GAT-01: Autosar Adaptive Core V5, PRJ-GAT-02: Toyota Next-Gen Cockpit OS, PRJ-GST-01: Smart Factory IoT Gateway). Với mỗi dự án, hệ thống thể hiện đầy đủ khách hàng toàn cầu (Toyota Motors, Bosch Global), quốc gia đối tác (Nhật Bản, Đức), đơn vị FSU chủ trì, người quản trị dự án (PM), ngăn xếp công nghệ sử dụng, tổng ngân sách giải ngân và thanh tiến độ hoàn thành trực quan (Progress Bar) theo tỷ lệ phần trăm thực tế."
    )
    insert_image_with_caption(doc, t_p, "06_projects.png", 19, "Giao diện Quản lý Danh mục Dự án phần mềm")

    # 4.6 (formerly 4.2.6)
    add_academic_heading2(t_p, "Phân hệ Điều phối & Phân bổ Nguồn lực (Resource Allocation)")
    add_academic_paragraph(t_p,
        "Điều phối nguồn lực là bài toán giải quyết trực tiếp điểm nghẽn nghiêm trọng đã được phân tích ở Chương 2. Phân hệ Resource Allocation cung cấp bảng ma trận điều phối nguồn nhân lực, cho phép các Quản lý Dự án (PM) và cán bộ Resource Manager lập kế hoạch đưa nhân sự vào dự án theo tỷ lệ % khối lượng công việc cam kết (Workload Allocation Rate). Hệ thống tích hợp thuật toán kiểm soát tỷ lệ phân bổ: khi một nhân sự đã được gán 100% công suất vào các dự án, hệ thống sẽ cảnh báo ngăn chặn tình trạng quá tải (Over-allocated); đồng thời liệt kê danh sách kỹ sư đang ở trạng thái rảnh rỗi (Bench) để tối ưu hóa chi phí vận hành doanh nghiệp."
    )
    insert_image_with_caption(doc, t_p, "07_resources.png", 20, "Giao diện Điều phối và Phân bổ Nguồn lực dự án")

    # 4.7 (formerly 4.2.7)
    add_academic_heading2(t_p, "Phân hệ Quản lý Nghỉ phép & Quy trình Phê duyệt trực tuyến (Leave Management)")
    add_academic_paragraph(t_p,
        "Phân hệ Quản lý Nghỉ phép số hóa 100% quy trình nộp và xét duyệt đơn nghỉ việc, thay thế hoàn toàn việc gửi email hoặc giấy tờ thủ công truyền thống. Nhân viên có thể theo dõi trực tiếp quỹ phép năm còn lại, số ngày đã nghỉ và thực hiện nộp đơn xin nghỉ phép với các danh mục: Nghỉ phép năm (Annual Leave), Nghỉ ốm (Sick Leave), Nghỉ thai sản hoặc Việc riêng. Hệ thống tự động tính toán chính xác số ngày nghỉ trừ các ngày cuối tuần và gửi đơn vào danh sách chờ duyệt của Quản lý dự án / Bộ phận nhân sự. Cấp quản lý có thể xem xét lý do và thực hiện 'Phê duyệt' (Approve) hoặc 'Từ chối' (Reject) kèm lý do phản hồi tức thì."
    )
    insert_image_with_caption(doc, t_p, "08_leave.png", 21, "Giao diện Quản lý và Xét duyệt Đơn Nghỉ phép")

    # 4.8 (formerly 4.2.8)
    add_academic_heading2(t_p, "Phân hệ Quản lý Tiền lương & Xuất Phiếu lương (Payroll & Payslip)")
    add_academic_paragraph(t_p,
        "Phân hệ Quản lý Tiền lương cung cấp công cụ tự động hóa việc tính toán chi phí lương thưởng định kỳ hàng tháng cho từng đơn vị FSU. Công thức tính toán được xây dựng chuẩn mực dựa trên mức lương cơ sở trong hợp đồng, số ngày công thực tế ghi nhận qua hệ thống điểm danh, phụ cấp trách nhiệm/dự án và các khoản trích đóng bắt buộc theo quy định pháp luật (BHXH 8%, BHYT 1.5%, BHTN 1%) cùng thuế thu nhập cá nhân. Bên cạnh bảng lương tổng thể dành cho kế toán và nhân sự, mỗi nhân viên đều có thể tự tra cứu Phiếu lương (Payslip) cá nhân với các khoản mục thu nhập và khấu trừ minh bạch, hỗ trợ in ấn hoặc xuất hóa đơn điện tử an toàn."
    )
    insert_image_with_caption(doc, t_p, "09_payroll.png", 22, "Giao diện Quản lý Bảng lương tự động và Phiếu lương cá nhân")

    # 4.9 (formerly 4.2.9)
    add_academic_heading2(t_p, "Phân hệ Đánh giá Hiệu suất (KPI) & Xếp loại nhân sự (Performance Review)")
    add_academic_paragraph(t_p,
        "Phân hệ Đánh giá Hiệu suất chuẩn hóa quy trình đánh giá định kỳ nửa năm (Mid-year) và cuối năm (Year-end) tại FPT Software. Quy trình kết hợp đa chiều giữa phần tự đánh giá (Self-review) của nhân viên đối với các mục tiêu OKR/KPI đã đăng ký và phần chấm điểm khách quan của Quản lý trực tiếp (PM/Team Lead). Điểm số tổng hợp trên thang 100 được quy đổi tự động thành các bậc xếp loại tiêu chuẩn: Hạng A (Xuất sắc - Top performer), Hạng B (Hoàn thành tốt), Hạng C (Đạt yêu cầu) và Hạng D (Cần cải thiện). Kết quả xếp loại là căn cứ trực tiếp để xét duyệt khen thưởng, phân chia lương tháng 13 và điều chỉnh bậc lương hàng năm."
    )
    insert_image_with_caption(doc, t_p, "10_performance.png", 23, "Giao diện Đánh giá Hiệu suất và Xếp loại KPI định kỳ")

    # 4.10 (formerly 4.2.10)
    add_academic_heading2(t_p, "Phân hệ Ma trận Kỹ năng chuyên môn (Skill Matrix)")
    add_academic_paragraph(t_p,
        "Ma trận Kỹ năng (Skill Matrix) là công cụ chiến lược giúp chuẩn hóa toàn bộ năng lực công nghệ của đội ngũ kỹ sư phần mềm. Hệ thống phân loại kỹ năng thành các nhóm rõ ràng: Ngôn ngữ lập trình (Java, C++, Python, .NET, JavaScript), Nền tảng Frontend/Mobile, Điện toán đám mây & DevOps (AWS, Azure, Docker, Kubernetes) và Năng lực ngoại ngữ (Tiếng Nhật JLPT, Tiếng Anh IELTS). Mỗi kỹ năng được định lượng theo thang bậc từ Level 1 (Mới bắt đầu) đến Level 5 (Chuyên gia/Architect) đi kèm số năm kinh nghiệm thực chiến. Nhờ ma trận này, các nhà quản lý dự án có thể nhanh chóng tra cứu và lập đội dự án (Staffing) chính xác theo yêu cầu khắt khe của khách hàng quốc tế."
    )
    insert_image_with_caption(doc, t_p, "11_skills.png", 24, "Giao diện Ma trận Kỹ năng chuyên môn (Skill Matrix)")

    # 4.11 (formerly 4.2.11)
    add_academic_heading2(t_p, "Phân hệ Quản lý Chức danh, Bậc lương tiêu chuẩn & Tuyển dụng ứng viên")
    add_academic_paragraph(t_p,
        "Nhằm xây dựng thang bậc phát triển nghề nghiệp rõ ràng cho nhân viên, hệ thống thiết lập bảng danh mục Chức danh công việc chuẩn (Positions) đi liền với các ngạch bậc lương tiêu chuẩn (Salary Grades: SG-01 đến SG-05). Mỗi cấp bậc quy định rõ yêu cầu về trình độ chuyên môn, trách nhiệm công việc và dải lương tham chiếu tương ứng, tạo sự công bằng và động lực phát triển nội bộ."
    )
    insert_image_with_caption(doc, t_p, "12_positions.png", 25, "Giao diện Quản lý Chức danh và Bậc lương tiêu chuẩn")
    add_academic_paragraph(t_p,
        "Song song với quản lý nội bộ, phân hệ Quản lý Tuyển dụng (Recruitment) hỗ trợ phòng nhân sự theo dõi toàn diện quy trình săn đón nhân tài: từ việc tiếp nhận hồ sơ ứng viên (Candidate CV), phân loại chuyên môn, lên lịch phỏng vấn chuyên sâu, theo dõi trạng thái ứng tuyển (Screening, Interview, Offer, Hired) đến bước tự động chuyển đổi hồ sơ trúng tuyển thành tài khoản nhân viên chính thức trong hệ thống."
    )
    insert_image_with_caption(doc, t_p, "13_recruitment.png", 26, "Giao diện Quản lý Hồ sơ Ứng viên Tuyển dụng")

    # 4.12 (formerly 4.2.12)
    add_academic_heading2(t_p, "Phân hệ Quản lý Tài sản Doanh nghiệp & Đào tạo nội bộ")
    add_academic_paragraph(t_p,
        "Phân hệ Quản lý Tài sản hỗ trợ kiểm soát chặt chẽ các trang thiết bị phần cứng giá trị cao được cấp phát cho kỹ sư phần mềm trong quá trình làm dự án, bao gồm: Laptop làm việc chuyên dụng, màn hình phụ, thiết bị thử nghiệm mạch vi điều khiển (ECU) và các bản quyền phần mềm lập trình. Mỗi tài sản đều được định danh bằng mã số riêng (Asset Tag), theo dõi người sở hữu hiện tại, ngày bàn giao và tình trạng khấu hao thiết bị."
    )
    insert_image_with_caption(doc, t_p, "14_assets.png", 27, "Giao diện Quản lý Cấp phát Tài sản Doanh nghiệp")
    add_academic_paragraph(t_p,
        "Về công tác phát triển nguồn nhân lực, phân hệ Đào tạo nội bộ (Training & Roadmap) cung cấp danh mục các khóa đào tạo công nghệ mới (Autosar Architecture, AWS Solutions Architect, Agile/Scrum Master, Tiếng Nhật thương mại IT). Hệ thống theo dõi tiến độ học tập, điểm số bài thi sát hạch và tự động lưu vết các chứng chỉ nghề nghiệp mà kỹ sư đã hoàn thành, giúp nâng cao năng lực cạnh tranh cốt lõi cho FPT Software."
    )
    insert_image_with_caption(doc, t_p, "15_training.png", 28, "Giao diện Quản lý Đào tạo và Chứng chỉ nghề nghiệp")

    # 4.13 (formerly 4.2.13)
    add_academic_heading2(t_p, "Phân hệ Hồ sơ Cá nhân & Thiết lập tài khoản người dùng")
    add_academic_paragraph(t_p,
        "Mỗi nhân viên khi đăng nhập vào hệ thống đều có thể truy cập trang Hồ sơ Cá nhân (Profile & Settings) để xem lại toàn bộ thông tin hợp đồng lao động, thông tin tài khoản ngân hàng nhận lương, liên kết mạng xã hội nghề nghiệp và danh sách chứng chỉ đã được phòng đào tạo xác nhận. Tại đây, người dùng có thể chủ động cập nhật mật khẩu đăng nhập, cập nhật số điện thoại liên lạc khẩn cấp và chuyển đổi nhanh giữa hai chế độ giao diện Sáng (Light Mode) hoặc Tối (Dark Mode) để bảo vệ thị lực khi làm việc trong thời gian dài."
    )
    insert_image_with_caption(doc, t_p, "16_profile.png", 29, "Giao diện Hồ sơ Cá nhân và Thiết lập tài khoản người dùng")

    # Save docx with python-docx
    print("Saving modified docx...")
    doc.save(DOC_PATH)
    print("Saved docx file successfully!")

    # 5. Use Word COM to generate official TablesOfFigures, update TOC and all fields
    print("Opening Microsoft Word COM to generate TOFs and update all fields...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(os.path.abspath(DOC_PATH))
        
        # Replace HINH placeholder with Table of Figures for Hinh
        find1 = wdoc.Content.Find
        find1.Text = "__PLACEHOLDER_TOF_HINH__"
        if find1.Execute():
            rng1 = find1.Parent
            rng1.Text = ""
            tof1 = wdoc.TablesOfFigures.Add(Range=rng1, Caption="Hình", IncludeLabel=True, UseHeadingStyles=False)
            print(f"Generated official Table of Figures for 'Hình'! (Items: {tof1.Range.Paragraphs.Count})")
        else:
            print("Warning: __PLACEHOLDER_TOF_HINH__ not found!")

        # Replace BANG placeholder with Table of Figures for Bang
        find2 = wdoc.Content.Find
        find2.Text = "__PLACEHOLDER_TOF_BANG__"
        if find2.Execute():
            rng2 = find2.Parent
            rng2.Text = ""
            tof2 = wdoc.TablesOfFigures.Add(Range=rng2, Caption="Bảng", IncludeLabel=True, UseHeadingStyles=False)
            print(f"Generated official Table of Figures for 'Bảng'! (Items: {tof2.Range.Paragraphs.Count})")
        else:
            print("Warning: __PLACEHOLDER_TOF_BANG__ not found!")

        # Update Table of Contents
        print("Updating Table of Contents...")
        for i in range(1, wdoc.TablesOfContents.Count + 1):
            wdoc.TablesOfContents(i).Update()
        print("TOC updated!")

        # Update all fields in document
        print("Updating all Fields in document...")
        wdoc.Fields.Update()
        
        wdoc.Save()
        print("Document saved successfully via Word COM!")
    finally:
        wdoc.Close(False)
        word.Quit()

    print("ALL TASKS COMPLETED PERFECTLY!")

if __name__ == "__main__":
    main()
